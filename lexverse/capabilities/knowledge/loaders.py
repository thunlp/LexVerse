from __future__ import annotations

import csv
import hashlib
import json
from pathlib import Path
import sqlite3
import tempfile

from filelock import FileLock

from langchain_core.documents import Document

from lexverse.config import ConfigError
from lexverse.capabilities.knowledge.sources import MANIFEST_NAME, field_value, read_source_manifest


SUPPORTED = {".txt", ".md", ".json", ".jsonl", ".csv", ".pdf", ".docx"}
KINDS = {"legal_laws", "legal_cases", "legal_concepts", "legal_qa", "legal_templates"}
LOADER_VERSION = 5
ANNOTATION_FILES = {
    ("legal_cases", "lexchain/test_reference.json"),
    ("legal_cases", "legal_world/light_case_dataset.json"),
}
FIELDS = {
    "legal_laws": ["law_name", "article", "text", "effective_from", "effective_to"],
    "legal_concepts": ["term", "category", "domain", "concept", "entity_name", "facts"],
    "legal_templates": ["template_type", "title", "content", "case_type"],
    "legal_qa": ["question", "answer", "opening", "facts", "roles", "topic_list", "theme"],
    "legal_cases": ["caseID", "content", "first_instance", "second_instance", "qw", "fact", "reason",
                    "案件名", "法院", "案由", "案号", "发布时间", "基本案情", "法院意见", "文书",
                    "case_r_f", "laws", "claim", "evidence", "text", "plaintiff_claim",
                    "plaintiff_case_details", "defendant_defence", "other_statement", "court_information",
                    "CaseId", "Fact", "Reasoning", "一审法院认定事实", "一审法院意见", "本院二审查明事实", "二审意见",
                    "案件描述", "法律文书原文", "defendant", "lawyer", "procurator"],
}
EXCLUDED = {"gold", "label", "labels", "split_labels", "rubric", "score", "reference",
            "ground_truth", "reference_answer", "evaluation_criteria"}
CASE_EXCLUDED = EXCLUDED | {"result", "charge", "article"}


def hash_file(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        for block in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def scan_source(root: Path) -> tuple[dict[str, str], list[dict]]:
    files, skipped = {}, []
    for path in sorted(root.rglob("*")):
        relative = path.relative_to(root).as_posix()
        if path.is_symlink():
            raise ConfigError(f"Knowledge symbolic link is not allowed: {relative}")
        if not path.is_file():
            continue
        if relative == MANIFEST_NAME:
            files[relative] = hash_file(path)
        elif any(part.startswith(".") or part.startswith("~$") for part in path.relative_to(root).parts):
            skipped.append({"file": relative, "reason": "hidden_or_temporary"})
        elif _annotation_file(path):
            skipped.append({"file": relative, "reason": "benchmark_annotation"})
        elif path.suffix.lower() not in SUPPORTED:
            skipped.append({"file": relative, "reason": "unsupported_format"})
        else:
            files[relative] = hash_file(path)
    return files, skipped


def source_kind(root: Path) -> str:
    return next((part for part in reversed(root.parts) if part in KINDS), "user")


def _annotation_file(path: Path) -> bool:
    for parent in path.parents:
        if parent.name in KINDS:
            return (parent.name, path.relative_to(parent).as_posix()) in ANNOTATION_FILES
    return False


def _parsed_records(path: Path, kind: str, content_hash: str, cache: Path | None, mapping=None):
    parse = lambda: _parse(path, kind, mapping) if mapping is not None else _parse(path, kind)
    if cache is None:
        yield from parse()
        return
    identity = [LOADER_VERSION, kind, path.suffix.lower(), path.stem, content_hash,
                mapping.model_dump() if mapping is not None else None]
    key = hashlib.sha256(json.dumps(identity, ensure_ascii=False).encode()).hexdigest()
    destination = cache / "parsed_documents" / (key + ".sqlite")
    destination.parent.mkdir(parents=True, exist_ok=True)
    with FileLock(str(destination) + ".lock"):
        if not destination.exists():
            with tempfile.TemporaryDirectory(dir=destination.parent) as temporary:
                staged = Path(temporary) / "records.sqlite"
                db = sqlite3.connect(staged)
                try:
                    db.execute("CREATE TABLE records (ordinal INTEGER PRIMARY KEY, payload TEXT, checksum TEXT)")
                    db.execute("CREATE TABLE manifest (record_count INTEGER NOT NULL)")
                    count = 0
                    for ordinal, record in enumerate(parse()):
                        count += 1
                        if not record[1].strip():
                            raise ValueError("No extractable text")
                        payload = json.dumps(record, ensure_ascii=False)
                        db.execute("INSERT INTO records VALUES (?, ?, ?)",
                                   (ordinal, payload, hashlib.sha256(payload.encode()).hexdigest()))
                    if hash_file(path) != content_hash:
                        raise ValueError("Source changed during parsing")
                    if not count:
                        raise ValueError("Source contains no records")
                    db.execute("INSERT INTO manifest VALUES (?)", (count,))
                    db.commit()
                finally:
                    db.close()
                staged.rename(destination)
    db = sqlite3.connect(destination.resolve().as_uri() + "?mode=ro", uri=True)
    try:
        expected = db.execute("SELECT record_count FROM manifest").fetchone()
        if expected is None or db.execute("SELECT count(*) FROM records").fetchone()[0] != expected[0]:
            raise ValueError("Parsed document cache is incomplete")
        for payload, checksum in db.execute("SELECT payload, checksum FROM records ORDER BY ordinal"):
            if hashlib.sha256(payload.encode()).hexdigest() != checksum:
                raise ValueError("Parsed document cache is damaged")
            yield json.loads(payload)
    finally:
        db.close()


def load_documents(root: Path, files: dict[str, str], source_id: str, *, cache: Path | None = None, use_manifest=True, file_completed=None):
    mapping = read_source_manifest(root, files.get(MANIFEST_NAME)) if use_manifest else None
    if mapping is not None and (MANIFEST_NAME not in files or hash_file(root / MANIFEST_NAME) != files[MANIFEST_NAME]):
        raise ConfigError("Knowledge manifest changed after scanning")
    for relative, content_hash in files.items():
        if relative == MANIFEST_NAME:
            if file_completed is not None:
                file_completed(relative)
            continue
        path = root / relative
        kind = mapping.kind if mapping is not None else source_kind(path.parent)
        try:
            for position, text, title, extra in _parsed_records(path, kind, content_hash, cache, mapping):
                if not text.strip():
                    raise ValueError("No extractable text")
                key = f"{source_id}:{relative}:{position}"
                yield Document(page_content=text, metadata={
                    "source_id": source_id, "document_id": hashlib.sha256(key.encode()).hexdigest(),
                    "kind": kind, "title": title or path.stem, "source_path": relative,
                    "position": position, "content_hash": content_hash, **extra,
                })
            if hash_file(path) != content_hash:
                raise ValueError("Source changed during parsing")
        except Exception as exc:
            raise ConfigError(f"Knowledge parsing failed: {relative} ({type(exc).__name__}: {exc})") from exc
        if file_completed is not None:
            file_completed(relative)
    if use_manifest:
        manifest_path = root / MANIFEST_NAME
        if manifest_path.is_symlink():
            raise ConfigError("Knowledge manifest must not be a symbolic link")
        try:
            current_hash = hash_file(manifest_path) if manifest_path.exists() else None
        except OSError as exc:
            raise ConfigError("Cannot verify knowledge manifest after parsing") from exc
        if current_hash != files.get(MANIFEST_NAME):
            raise ConfigError("Knowledge manifest changed during parsing")


def _parse(path: Path, kind: str, mapping=None):
    suffix = path.suffix.lower()
    if suffix in {".txt", ".md"}:
        yield "text", path.read_text(encoding="utf-8-sig"), path.stem, {}
    elif suffix == ".pdf":
        from langchain_community.document_loaders import PyPDFLoader
        pages = PyPDFLoader(str(path)).load()
        if not pages:
            raise ValueError("PDF has no pages")
        if not any(page.page_content.strip() for page in pages):
            raise ValueError("PDF has no extractable text; OCR is not supported")
        for number, page in enumerate(pages, 1):
            if not page.page_content.strip():
                continue
            yield f"page:{number}", page.page_content, path.stem, {"page": number}
    elif suffix == ".docx":
        from docx import Document as WordDocument
        doc = WordDocument(path)
        lines = [p.text for p in doc.paragraphs if p.text.strip()]
        lines.extend(" | ".join(cell.text for cell in row.cells) for table in doc.tables for row in table.rows)
        yield "text", "\n".join(lines), path.stem, {}
    else:
        count = 0
        for index, row in _records(path, "user" if mapping is not None else kind):
            count += 1
            title, extra = path.stem, {}
            if mapping is not None:
                excluded = set(mapping.exclude_fields) | (CASE_EXCLUDED if kind == "legal_cases" else EXCLUDED if kind != "user" else set())
                mapped_fields = [*mapping.text_fields, *mapping.metadata_fields.values()]
                if mapping.title_field:
                    mapped_fields.append(mapping.title_field)
                if any(part in excluded for field in mapped_fields for part in field.split(".")):
                    raise ValueError("A mapped text field is excluded")
                text = _text({field: field_value(row, field) for field in mapping.text_fields}, exclude=excluded)
                title = str(field_value(row, mapping.title_field)) if mapping.title_field else path.stem
                extra = {name: str(field_value(row, field)) for name, field in mapping.metadata_fields.items()}
            elif isinstance(row, dict) and kind != "user":
                if not any(key in row for key in FIELDS[kind]):
                    raise ValueError("Unknown legal record schema; provide a reviewed field mapping")
                chosen = {key: row[key] for key in FIELDS[kind] if key in row}
                title = str(row.get("law_name") or row.get("title") or row.get("term") or row.get("caseID") or row.get("CaseId") or row.get("案号") or path.stem)
                extra = {key: row[key] for key in ("article", "effective_from", "effective_to", "source") if key in row}
                text = _text(chosen, exclude=CASE_EXCLUDED if kind == "legal_cases" else EXCLUDED)
            else:
                text = _text(row)
            yield f"row:{index}", text, title, extra
        if not count:
            raise ValueError("Source contains no records")


def _records(path: Path, kind: str):
    if path.suffix.lower() == ".csv":
        with path.open(encoding="utf-8-sig", newline="") as stream:
            yield from enumerate(csv.DictReader(stream), 1)
    elif path.suffix.lower() == ".jsonl":
        with path.open(encoding="utf-8-sig") as stream:
            for line_number, line in enumerate(stream, 1):
                if line.strip():
                    yield line_number, json.loads(line)
    else:
        value = json.loads(path.read_text(encoding="utf-8-sig"))
        rows = value if isinstance(value, list) else (
            value.values() if kind != "user" and isinstance(value, dict) and value
            and not any(key in value for key in FIELDS[kind])
            and all(isinstance(v, dict) for v in value.values()) else [value]
        )
        yield from enumerate(rows, 1)


def _text(value, *, exclude=frozenset()) -> str:
    if isinstance(value, dict):
        return "\n".join(f"{key}: {_text(item, exclude=exclude)}" for key, item in value.items() if key not in exclude)
    if isinstance(value, list):
        return "\n".join(_text(item, exclude=exclude) for item in value)
    return "" if value is None else str(value)
